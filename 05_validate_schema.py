from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from schema_contract import (
    PROJECT_ROOT,
    SCHEMA,
    SCHEMA_PATH,
    normalize_header,
    validate_archive,
    validate_header_archive,
)


DOC_DIR = PROJECT_ROOT / "data" / "documentation"
LAYOUT_PATH = DOC_DIR / "file_layout_july_2026.xlsx"
REPORT_PATH = PROJECT_ROOT / "data" / "inventory" / "release_47_schema_validation.json"


def validate_workbook() -> dict:
    if not LAYOUT_PATH.exists():
        raise FileNotFoundError(f"Official Release 47 file layout not found: {LAYOUT_PATH}")
    workbook = load_workbook(LAYOUT_PATH, read_only=True, data_only=True)
    sheet_names = {
        "origination": "Origination Data File",
        "performance": "Monthly Performance Data File",
    }
    results = {}
    for file_type, sheet_name in sheet_names.items():
        sheet = workbook[sheet_name]
        rows = list(sheet.iter_rows(min_row=3, values_only=True))
        fields = [(int(row[0]), str(row[1])) for row in rows if row[0] is not None]
        expected = SCHEMA[file_type]
        expected_positions = list(range(1, len(expected) + 1))
        observed_positions = [position for position, _ in fields]
        if observed_positions != expected_positions:
            raise ValueError(
                f"Workbook positions are not continuous for {file_type}: {observed_positions}"
            )
        if len(fields) != len(expected):
            raise ValueError(
                f"Workbook field count mismatch for {file_type}: "
                f"{len(fields)} observed, {len(expected)} expected."
            )

        workbook_header_names = [normalize_header(name) for _, name in fields]
        official_header_names = [
            normalize_header(name) for name in SCHEMA["header_names"][file_type]
        ]
        if workbook_header_names != official_header_names:
            differences = [
                (index + 1, workbook_name, header_name)
                for index, (workbook_name, header_name) in enumerate(
                    zip(workbook_header_names, official_header_names)
                )
                if workbook_name != header_name
            ]
            raise ValueError(
                f"Workbook fields do not match the Release 47 header for {file_type}: "
                f"{differences[:10]}"
            )
        if [field["position"] for field in expected] != expected_positions:
            raise ValueError(f"Versioned map positions do not match workbook for {file_type}.")
        results[file_type] = {"field_count": len(fields), "positions_verified": True}
    return results


def annual_sample_archives() -> list[Path]:
    by_name = {}
    for root in (PROJECT_ROOT, PROJECT_ROOT / "data" / "raw_zips"):
        if root.exists():
            for path in root.glob("sample_*.zip"):
                existing = by_name.get(path.name)
                if existing and existing.stat().st_size != path.stat().st_size:
                    raise ValueError(f"Conflicting copies of {path.name} have different sizes.")
                by_name[path.name] = path
    return sorted(by_name.values())


def main() -> None:
    workbook_result = validate_workbook()
    header_result = validate_header_archive()
    archives = annual_sample_archives()
    if not archives:
        raise FileNotFoundError("No annual sample_YYYY.zip files found.")
    archive_results = [validate_archive(path, full_scan=True) for path in archives]
    years = [int(path.stem.split("_")[-1]) for path in archives]
    report = {
        "schema_version": SCHEMA["schema_version"],
        "release": SCHEMA["release"],
        "release_date": SCHEMA["release_date"],
        "workbook_validation": workbook_result,
        "header_archive_validation": {
            file_type: len(fields) for file_type, fields in header_result.items()
        },
        "vintages": years,
        "archives": archive_results,
        "row_width_checks": "all nonempty text rows scanned and validated",
    }
    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    REPORT_PATH.write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(json.dumps(report, indent=2))
    print(f"Saved Release 47 schema validation report: {REPORT_PATH}")


if __name__ == "__main__":
    main()

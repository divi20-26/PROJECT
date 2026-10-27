from __future__ import annotations

import unittest
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from schema_contract import (
    PROJECT_ROOT,
    SCHEMA,
    field_names,
    normalize_header,
    validate_header_archive,
    validate_key_fields,
    validate_row,
)


class Release47SchemaTests(unittest.TestCase):
    def test_release_47_field_counts_and_contiguous_positions(self) -> None:
        self.assertEqual(len(SCHEMA["origination"]), 31)
        self.assertEqual(len(SCHEMA["performance"]), 35)
        for file_type in ("origination", "performance"):
            positions = [field["position"] for field in SCHEMA[file_type]]
            self.assertEqual(positions, list(range(1, len(positions) + 1)))

    def test_release_47_mapping_matches_official_workbook_and_headers(self) -> None:
        workbook = load_workbook(
            PROJECT_ROOT / "data" / "documentation" / "file_layout_july_2026.xlsx",
            read_only=True,
            data_only=True,
        )
        header_archive = validate_header_archive()
        sheet_map = {
            "origination": "Origination Data File",
            "performance": "Monthly Performance Data File",
        }
        for file_type, sheet_name in sheet_map.items():
            rows = [
                (int(row[0]), normalize_header(str(row[1])))
                for row in workbook[sheet_name].iter_rows(min_row=3, values_only=True)
                if row[0] is not None
            ]
            actual_headers = [
                normalize_header(value) for value in header_archive[file_type]
            ]
            self.assertEqual([position for position, _ in rows], list(range(1, len(rows) + 1)))
            self.assertEqual([name for _, name in rows], actual_headers)
            self.assertEqual(len(rows), len(field_names(file_type)))

    def test_critical_field_positions(self) -> None:
        self.assertEqual(SCHEMA["origination"][19]["name"], "loan_identifier")
        self.assertEqual(SCHEMA["origination"][1]["name"], "first_payment_date")
        self.assertEqual(SCHEMA["performance"][0]["name"], "loan_identifier")
        self.assertEqual(SCHEMA["performance"][1]["name"], "period")
        self.assertEqual(SCHEMA["performance"][3]["name"], "current_loan_delinquency_status")
        self.assertEqual(SCHEMA["performance"][8]["name"], "zero_balance_code")
        self.assertEqual(SCHEMA["performance"][9]["name"], "zero_balance_effective_date")

    def test_row_count_mismatch_fails(self) -> None:
        with self.assertRaisesRegex(ValueError, "Schema mismatch"):
            validate_row("origination", "A|B", "bad-file")

    def test_loan_id_and_reporting_period_positions_are_checked(self) -> None:
        valid_orig = [""] * 31
        valid_orig[1] = "200503"
        valid_orig[19] = "F05Q10000006"
        validate_key_fields("origination", valid_orig, "test-orig")

        valid_perf = [""] * 35
        valid_perf[0] = "F05Q10000006"
        valid_perf[1] = "200502"
        validate_key_fields("performance", valid_perf, "test-perf")

        valid_perf[1] = "BAD"
        with self.assertRaisesRegex(ValueError, "Invalid performance period"):
            validate_key_fields("performance", valid_perf, "test-perf")

    def test_all_downloaded_annual_vintage_samples_match_release_47(self) -> None:
        roots = [PROJECT_ROOT, PROJECT_ROOT / "data" / "raw_zips"]
        archives_by_name = {}
        for root in roots:
            if root.exists():
                for path in root.glob("sample_*.zip"):
                    previous = archives_by_name.get(path.name)
                    if previous:
                        self.assertEqual(previous.stat().st_size, path.stat().st_size)
                    archives_by_name[path.name] = path
        archives = sorted(archives_by_name.values())
        self.assertTrue(archives, "Expected downloaded annual Freddie Mac sample ZIP files.")
        years = set()
        for path in archives:
            year = int(path.stem.split("_")[-1])
            years.add(year)
            with zipfile.ZipFile(path) as archive:
                members = {
                    "origination": next(
                        name for name in archive.namelist()
                        if "orig" in Path(name).name.lower() and name.lower().endswith(".txt")
                    ),
                    "performance": next(
                        name for name in archive.namelist()
                        if "perf" in Path(name).name.lower() and name.lower().endswith(".txt")
                    ),
                }
                for file_type, member in members.items():
                    with archive.open(member) as stream:
                        rows = 0
                        for line_number, raw_line in enumerate(stream, start=1):
                            if not raw_line.strip():
                                continue
                            fields = validate_row(
                                file_type,
                                raw_line.decode("latin1"),
                                f"{path.name}:{member}:{line_number}",
                            )
                            if rows == 0:
                                validate_key_fields(
                                    file_type, fields, f"{path.name}:{member}:{line_number}"
                                )
                            rows += 1
                        self.assertGreater(rows, 0, f"Empty file: {path.name}:{member}")
        self.assertGreaterEqual(min(years), 2005)
        self.assertGreaterEqual(max(years), 2026)


if __name__ == "__main__":
    unittest.main()
